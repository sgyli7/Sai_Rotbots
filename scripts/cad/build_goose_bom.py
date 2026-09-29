"""Build an explicitly provisional Goose RC2 BOM from current CAD and spec.

Public USD unit prices are dated reference values, never China order quotes.
The preserved handoff workbook is deliberately not overwritten.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from sai_agent.paths import resource_root


def main() -> None:
    robot = resource_root() / 'robots/Goose_V0.1'
    spec_path = robot / 'configs/robot_spec.json'
    spec = json.loads(spec_path.read_text())
    manifest_path = robot / 'cad/exports/cad_manifest.json'
    cad = json.loads(manifest_path.read_text())
    spec_hash = hashlib.sha256(spec_path.read_bytes()).hexdigest()
    if cad['spec_sha256'] != spec_hash:
        raise ValueError('CAD and specification revisions differ')
    counts = {key: sum(j['servo'] == key for j in spec['joints']) for key in spec['servos']}
    rows = []

    def item(group, description, sku, quantity, price, source, gate, note='', unit='pcs'):
        rows.append(dict(group=group, description=description, sku=sku,
                         quantity=quantity, unit=unit, usd_public_unit_reference=price,
                         usd_public_extended_reference=round(quantity * price, 2) if price is not None else '',
                         china_quote_rmb='', source=source, gate=gate, note=note))

    item('actuator', 'XM430-W350-T', '902-0124-000', counts['xm430'], 269.90,
         'https://en.robotis.com/shop_en/item.php?it_id=902-0124-000', 'china_quote_and_TTL_pack_check',
         'Front HN12 horn and one ordinary X3P cable included per motor.')
    item('actuator', 'XM540-W270-T', '902-0137-000', counts['xm540'], 419.90,
         'https://en.robotis.com/shop_en/item.php?it_id=902-0137-000', 'china_quote_and_TTL_pack_check',
         'Front HN13 horn and one ordinary X3P cable included per motor.')
    item('actuator', 'XC330-M288-T', '902-0173-000', counts['xc330'], 89.90,
         'https://en.robotis.com/shop_en/item.php?it_id=902-0173-000', 'china_quote_and_mount_check',
         'Built-in front output, ordinary X3P cable included; no rear idler.')
    item('interface', 'U2D2 USB-C revision', '902-0132-000', 2, 32.10,
         'https://en.robotis.com/shop_en/item.php?it_id=902-0132-000', 'china_quote_and_USB_C_revision',
         'USB cable included; U2D2 does not power servos.')
    item('idler', 'HN12-I101 Set', '903-0240-000', counts['xm430'], 17.80,
         'https://en.robotis.com/shop_en/item.php?it_id=903-0240-000', 'conditional_full_rear_support',
         'Subtract sets supplied inside any selected FR12 frame kits.')
    item('idler', 'HN13-I101 Set', '903-0267-000', counts['xm540'], 42.90,
         'https://en.robotis.com/shop_en/item.php?it_id=903-0267-000', 'conditional_full_rear_support',
         'Subtract sets supplied inside any selected FR13 frame kits.')
    item('idler', 'FPX330-H101 4PCS SET', '903-0302-000', 1, 9.80,
         'https://en.robotis.com/shop_en/item.php?it_id=903-0302-000', 'conditional_full_rear_support',
         'One pack has four hinge/idler/cap groups; only two are used here.')
    other = [
        ('vision', 'OV5693 5MP USB Camera (A)', 'Waveshare 24710', 1, 'https://www.waveshare.com/ov5693-5mp-usb-camera-a.htm', 'physical_focus_FOV_mount_and_china_quote'),
        ('compute', 'ZERO 3W', 'Radxa ZERO 3W', 1, 'https://docs.radxa.com/zero/zero3', 'revision_connector_thermal_and_china_quote'),
        ('compute', 'USB HUB HAT (B)', 'Waveshare 20317', 1, 'https://www.waveshare.com/usb-hub-hat-b.htm', 'VBUS_isolation_and_total_port_power'),
        ('sensor', 'WT901C-TTL IMU', 'WitMotion WT901C-TTL', 1, 'https://www.wit-motion.com/', 'mount_orientation_and_timestamp_calibration'),
        ('sensor', 'USB-TTL adapter for IMU', 'WitMotion A.000057', 1, 'https://www.wit-motion.com/', 'revision_and_plug_envelope'),
        ('power', '5V D24V90F5 buck', 'Pololu 2866', 2, 'https://www.pololu.com/product/2866', 'servo_and_logic_independent_thermal_test'),
        ('audio', 'adjustable audio buck', 'Pololu 3798', 1, 'https://www.pololu.com/product/3798', 'audio_rail_thermal_and_wiring'),
        ('audio', 'DFPlayer Mini', 'DFRobot DFR0299', 1, 'https://www.dfrobot.com/product-1121.html', 'physical_speaker_loudness_and_UART'),
        ('audio', '30 mm speaker', 'Same Sky CMS-3058-18L200', 1, 'https://www.sameskydevices.com/', 'rear_mount_and_acoustic_test'),
        ('battery', '3S 1500 mAh XT60 LiPo', 'CNHL 1501303BK', 1, 'https://chinahobbyline.com/products/cnhl-black-series-v2-0-1500mah-11-1v-3s-130c-lipo-battery-with-xt60-plug', 'battery_protection_charge_and_access'),
        ('power', '12V motor rail relay', 'Panasonic CB1a-12V', 1, 'https://industry.panasonic.com/', 'DC_load_rating_and_EStop_test'),
        ('safety', 'NC latching E-stop control switch', 'Schneider XB5AS8442', 1, 'https://www.se.com/', 'external_control_box_and_relay_interface'),
        ('structure', '12x12x1.5 aluminium square tube, cut 134 mm', 'generic extrusion', 1, '', 'saw_drill_fixture_and_actual_stock'),
        ('wiring', '12V/5V TTL harness and branch power injection', 'custom', 1, '', 'lengths_AWG_connector_pinout_and_bend_sweep'),
        ('safety', 'main and branch fuses plus rated holders', 'select_by_current_curve', 1, '', 'current_curve_DC_breaking_and_wire_rating'),
        ('hardware', 'M2/M2.5/M3 structural fastener and insert schedule', 'select_after_stack_review', 1, '', 'diameter_length_thread_depth_and_access'),
        ('hardware', '6x2 mm service-door magnet', 'generic neodymium disc', 4, '', 'diameter_thickness_plating_polarity_and_retention'),
        ('hardware', 'service-door magnet retention adhesive', 'two-part epoxy small pack', 1, '', 'PETG_magnet_bond_and_access_test'),
    ]
    for group, description, sku, qty, source, gate in other:
        item(group, description, sku, qty, None, source, gate)
    material_masses={}
    for part in cad['parts']:
        key=(part['material'],part['color'])
        material_masses[key]=material_masses.get(key,0.)+part['mass_kg']
    for material,color,quantity,source in [
        ('petg','white',2.,'https://bambulab.cn/zh-cn/filament/petg-hf'),
        ('petg','orange',1.,'https://bambulab.cn/zh-cn/filament/petg-hf'),
        ('petg','black',1.,'https://bambulab.cn/zh-cn/filament/petg-hf'),
        ('tpu','black',1.,'https://bambulab.cn/zh-cn/filament/tpu-95a-hf'),
    ]:
        design_mass=material_masses.get((material,color),0.)
        item('print_material',f'{material.upper()} {color} stock',f'{material.upper()} {color}',quantity,None,
             source,'china_pack_price_drying_and_slicer_yield',
             f'CAD nominal {design_mass:.3f} kg; stock includes prototypes/support/waste; vendor pack size to confirm.',unit='kg')
    item('print_material','PAHT-CF structural trial stock','PAHT-CF',1.,None,
         'https://bambulab.cn/zh-cn/filament/paht-cf','conditional_strength_and_drying_trial',
         'Conditional sample stock only; not assumed in the 112 PETG parts or released structure.',unit='kg')

    hw = robot / 'hardware'
    fields = list(rows[0])
    with (hw / 'goose_rc2_bom.csv').open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields);writer.writeheader();writer.writerows(rows)

    wb = Workbook();ws = wb.active;ws.title = 'purchase_candidates'
    ws.append(fields)
    for row in rows:ws.append([row[key] for key in fields])
    print_ws = wb.create_sheet('print_and_cut_parts')
    print_fields = ['part', 'material', 'quantity', 'unit_mass_g', 'bounds_x_mm', 'bounds_y_mm',
                    'bounds_z_mm', 'x2d_single_10mm_margin_fit', 'source_step', 'source_stl', 'gate']
    print_ws.append(print_fields)
    for part in cad['parts']:
        box = part['bounds_body_mm']
        size = [round(box['max'][i] - box['min'][i], 2) for i in range(3)]
        ordered = sorted(size)
        fits = ordered[0] <= 236 and ordered[1] <= 236 and ordered[2] <= 240
        print_ws.append([part['name'], part['material'], 1, round(part['mass_kg'] * 1000, 2),
                         *size, fits, part['files']['step']['path'], part['files']['stl']['path'],
                         'slice_orientation_tolerance_strength_and_support' if part['material'] != 'aluminium'
                         else 'cut_drill_and_deburr'])
    facts = wb.create_sheet('read_me')
    for line in [
        ['status', 'Working engineering BOM; not a purchase release'],
        ['robot', spec['robot_id']], ['revision', spec['engineering_revision']],
        ['spec_sha256', spec_hash], ['cad_manifest_sha256', hashlib.sha256(manifest_path.read_bytes()).hexdigest()],
        ['public_price_basis', 'Global ROBOTIS USD listings checked 2026-09-28; no tax/shipping/conversion or China quote'],
        ['known_robotis_subtotal_usd', sum(x['usd_public_extended_reference'] or 0 for x in rows)],
        ['original_handoff_workbook', 'hardware/walker_r2_bom.xlsx remains untouched and historical'],
        ['printer_fit', 'AABB orientation fit only; final slicer/brim/support verification required'],
        ['power', 'No inferred battery/runtime/current/fuse rating from this sheet'],
    ]:facts.append(line)
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244352')
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(70,max(12,max(len(str(c.value or '')) for c in column)+2))
    wb.save(hw / 'goose_rc2_bom.xlsx')
    print(json.dumps({'rows':len(rows),'cad_parts':len(cad['parts']),
                      'known_robotis_usd_reference':round(sum(x['usd_public_extended_reference'] or 0 for x in rows),2),
                      'status':'working_not_order_ready'},indent=2))


if __name__ == '__main__':main()
