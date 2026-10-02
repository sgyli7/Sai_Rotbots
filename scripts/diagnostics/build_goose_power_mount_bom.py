"""Traceable incremental power-board mounting list, not a purchase release."""
from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
ROBOT = ROOT / 'robots/Goose_V0.1'


def main():
    config_path = ROBOT / 'configs/power_module_mounts.json'
    manifest_path = ROBOT / 'cad/exports/power_module_mounts/manifest.json'
    fit_path = ROBOT / 'evidence/power_module_mount_fit.json'
    config, kit, fit = [json.loads(p.read_text()) for p in [config_path, manifest_path, fit_path]]
    if not fit['named_sampled_fit_pass'] or not fit['bare_modules_in_original_boxes']:
        raise ValueError('Mounting list requires the finite nominal native fit screen')
    parts = {p['name']: p for p in kit['parts']}
    rows = []
    for module in config['modules']:
        name = module['name']
        row = dict(id=name + '_board', description=module['sku'], quantity=1,
            specification='5V bare module, actual OEM hole pattern; enclosed thermal/current release pending',
            mass_kg=module['mass_kg'], mass_basis='catalogue bare PCB mass, already contained in retained robot allocation',
            step='', source=module['sources'][0], domestic_price_rmb='', procurement_release=False,
            installation='Selected board, no wiring or thermal acceptance')
        rows.append(row)
        p = parts[name + '_frame_carrier']
        rows.append(dict(id=p['name'], description='Original CNC6061-T6 perimeter carrier', quantity=1,
            specification='2mm plate with integral M2 tapped posts; see native STEP and stack config',
            mass_kg=p['mass_kg'], mass_basis='native volume x assumed6061 density2700kg/m3',
            step=p['files']['step']['path'], source='configs/power_module_mounts.json', domestic_price_rmb='',
            procurement_release=False, installation='New carrier above current neck rear frame'))
    selectors = [
        ('pcb_insulating_washer', 'Nylon M2 insulating washer', 'OD4 / ID2.2 / t0.5mm',
         lambda n: '_pcb_' in n and n.endswith('_washer'), '12 new washers; material/preload and tolerance pending'),
        ('pcb_m2x6', 'Steel M2x6 socket screw candidate', 'nominal head OD3.8 / h2mm;6mm below bearing',
         lambda n: '_pcb_' in n and n.endswith('_m2x6'), 'Six new PCB bolts; M2 tap minor bore1.6mm'),
        ('frame_m3_washer', 'Steel M3 washer candidate', 'OD6 / ID3.2 / t0.5mm',
         lambda n: n.startswith('neck_power_frame_') and n.endswith('_washer'), 'Eight washers replace four existing two-washer stacks'),
        ('frame_m3x18', 'Steel M3x18 socket screw candidate', 'nominal head OD5.5 / h3mm;18mm below bearing',
         lambda n: n.startswith('neck_power_frame_') and n.endswith('_m3x18'), 'Four replace M3x16; new15.4mm stack,2.6mm nominal protrusion'),
        ('frame_m3_nut', 'Steel M3 nut candidate', 'across flats5.5 / h2.4mm; nominal tapped bore',
         lambda n: n.startswith('neck_power_frame_') and n.endswith('_m3_nut'), 'Four replace existing frame nuts; locking/preload remains pending'),
    ]
    covered = {row['id'] for row in rows if row['id'] in parts}
    for identifier, description, specification, predicate, installation in selectors:
        chosen = [p for name, p in parts.items() if predicate(name)]
        if not chosen:
            raise ValueError(('Empty hardware group', identifier))
        covered.update(p['name'] for p in chosen)
        rows.append(dict(id=identifier, description=description, quantity=len(chosen), specification=specification,
            mass_kg=sum(p['mass_kg'] for p in chosen), mass_basis='sum of nominal native envelopes, not measured helical hardware',
            step=chosen[0]['files']['step']['path'], source='configs/power_module_mounts.json',
            domestic_price_rmb='', procurement_release=False, installation=installation))
    if covered != set(parts):
        raise ValueError(('BOM native coverage mismatch', set(parts) - covered))
    directory = ROBOT / 'hardware'
    csv_path, xlsx_path = [directory / ('power_module_mount_bom.' + extension) for extension in ['csv', 'xlsx']]
    columns = list(rows[0])
    with csv_path.open('w', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=columns, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'power_mount_candidate'
    sheet.append(columns)
    for row in rows:
        sheet.append([row[column] for column in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244B5A')
    for index, column in enumerate(columns, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = 42 if column != 'quantity' else 12
    workbook.save(xlsx_path)
    inputs = [config_path, manifest_path, fit_path, Path(__file__)]
    result = dict(schema='goose_power_mount_incremental_bom_v1', rows=rows,
        native_part_coverage=len(covered), quantity_total=sum(row['quantity'] for row in rows),
        listed_bare_boards_and_native_hardware_mass_kg=sum(row['mass_kg'] for row in rows),
        conservative_robot_mass_increment_kg=fit['conservative_candidate_added_mass_kg'],
        already_retained_buck_allocations_kg=.065, full_robot_bom=False, purchase_release=False,
        source_hashes={str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        outputs={str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [csv_path, xlsx_path]},
        limitations=['Domestic pricing, exact fastener vendor, machining quote, cooling and harness are not closed.',
            'The36 native parts include replacement frame stacks; listing their mass is not the net robot mass increase.',
            'No current robot reserves were spent or removed.'])
    (directory / 'power_module_mount_bom.json').write_text(json.dumps(result, indent=2) + '\n')
    print('POWER BOM', len(rows), 'rows', result['quantity_total'], 'items', len(covered), 'native parts', flush=True)


if __name__ == '__main__':
    main()
