"""Audit Goose handoff preservation, local links and current artifact hashes.

This checks packaging integrity only; it cannot approve manufacturing or behavior.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import unquote
import zipfile

from openpyxl import load_workbook

from sai_agent.paths import resource_root


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def local_target(source: Path, raw: str) -> Path | None:
    raw = unquote(raw.strip('<>').split('#', 1)[0].split('?', 1)[0])
    if not raw or '://' in raw or raw.startswith(('mailto:', 'data:')):
        return None
    raw = re.sub(r':\d+$', '', raw)
    return Path(raw) if raw.startswith('/') else source.parent / raw


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, help='Write a new audit without overwriting the historical RC2 record')
    args = parser.parse_args()
    root = resource_root(); robot = root / 'robots/Goose_V0.1'
    archive = root / 'robots/Goose/goose_v0_1_workspace_handoff.zip'
    imported = robot / 'evidence/workspace_import_check.json'
    import_record = json.loads(imported.read_text())
    archive_matches = []
    changed_readmes = []
    missing_imports = []
    with zipfile.ZipFile(archive) as z:
        corrupt_member = z.testzip()
        members = [item for item in z.infolist() if not item.is_dir()]
        for item in members:
            if item.filename == 'README.md':
                path = robot / 'source/workspace_handoff_readme.md'
            else:
                path = root / item.filename
            if not path.exists():
                missing_imports.append(item.filename)
            elif hashlib.sha256(z.read(item)).hexdigest() == digest(path):
                archive_matches.append(item.filename)
            elif item.filename in ('README.md', 'robots/Goose_V0.1/README.md'):
                changed_readmes.append(item.filename)
            else:
                missing_imports.append(item.filename + ' (changed)')
        original_robot_readme = z.read('robots/Goose_V0.1/README.md').decode()
    historical_copy = (robot / 'source/robot_readme_handoff.md').read_text()
    expected_body = re.sub(r'(!?\[[^\]]*\]\()(?![a-z]+:|/|#)([^)]+)(\))',
                           lambda m: m.group(1) + '../' + m.group(2) + m.group(3),
                           original_robot_readme)
    historical_readme_preserved = historical_copy.endswith(expected_body)

    markdown_files = [root / 'README.md', root / 'docs/guides/goose_sai_lab_handoff.md'] + list(robot.rglob('*.md'))
    html_files = list(robot.rglob('*.html'))
    local_links = []
    broken = []
    for path in markdown_files:
        targets = re.findall(r'!?\[[^\]]*\]\(([^)]+)\)', path.read_text(errors='replace'))
        for raw in targets:
            target = local_target(path, raw)
            if target is not None:
                local_links.append((path, raw))
                if not target.exists():
                    broken.append({'source': str(path.relative_to(root)), 'target': raw})
    for path in html_files:
        targets = re.findall(r'\b(?:src|href)=["\']([^"\']+)["\']', path.read_text(errors='replace'))
        for raw in targets:
            target = local_target(path, raw)
            if target is not None:
                local_links.append((path, raw))
                if not target.exists():
                    broken.append({'source': str(path.relative_to(root)), 'target': raw})

    manifest_path = robot / 'cad/exports/cad_manifest.json'
    cad = json.loads(manifest_path.read_text())
    bad_cad = []
    for part in cad['parts']:
        for kind, file in part['files'].items():
            path = robot / file['path']
            if not path.exists() or digest(path) != file['sha256']:
                bad_cad.append(part['name'] + ':' + kind)
    spec_path = robot / 'configs/robot_spec.json'
    model_path = robot / 'models/full/robot.xml'
    transfer = json.loads((robot / 'models/full/rigid_transfer.json').read_text())
    bom_path = robot / 'hardware/goose_rc2_bom.xlsx'
    bom = load_workbook(bom_path, read_only=True, data_only=False)
    bom_csv = list(csv.DictReader((robot / 'hardware/goose_rc2_bom.csv').open(newline='')))
    workbook_rows = {sheet.title: sheet.max_row - 1 for sheet in bom}
    checks = {
        'original_zip_intact': corrupt_member is None and digest(archive) == import_record['source_archive_sha256'],
        'all_import_members_accounted_for': len(members) == import_record['archive_members'] and not missing_imports and len(archive_matches) + len(changed_readmes) == len(members),
        'readme_history_preserved': historical_readme_preserved and len(changed_readmes) == 2,
        'all_local_links_resolve': not broken,
        'all_cad_files_match_manifest': not bad_cad and len(cad['parts']) == 121 and cad['all_solids_valid'],
        'spec_model_transfer_match': cad['spec_sha256'] == digest(spec_path) and transfer['source_spec_sha256'] == digest(spec_path) and transfer['model_sha256'] == digest(model_path),
        'bom_and_cad_counts_match': len(bom_csv) == workbook_rows['purchase_candidates'] and workbook_rows['print_and_cut_parts'] == len(cad['parts']),
    }
    result = {
        'scope': 'Handoff and artifact integrity only; no manufacturing, walking, autonomy or real-hardware acceptance',
        'archive_sha256': digest(archive), 'archive_members': len(members),
        'archive_byte_identical_in_current_paths': len(archive_matches),
        'readmes_with_current_navigation_edits': changed_readmes,
        'historical_robot_readme_link_adjusted_copy_valid': historical_readme_preserved,
        'markdown_files_checked': len(markdown_files), 'html_files_checked': len(html_files),
        'local_links_checked': len(local_links), 'broken_links': broken,
        'cad_parts': len(cad['parts']), 'cad_files_checked': len(cad['parts']) * 3,
        'cad_hash_mismatches': bad_cad,
        'spec_sha256': digest(spec_path), 'model_sha256': digest(model_path),
        'transfer_sha256': digest(robot / 'models/full/rigid_transfer.json'),
        'bom_purchase_rows': len(bom_csv), 'bom_workbook_rows': workbook_rows,
        'checks': checks, 'pass': all(checks.values()),
    }
    output = args.output or robot / 'evidence/delivery_integrity_rc2.json'
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0 if result['pass'] else 2


if __name__ == '__main__':
    raise SystemExit(main())
