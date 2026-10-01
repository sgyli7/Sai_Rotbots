"""Delivery identity audit retaining OEM-fit and overload failures explicitly."""
from pathlib import Path
import csv
import hashlib
import json
import math
import re

import numpy as np
from PIL import Image
from openpyxl import load_workbook
from sai_agent.goose.mass_properties import aggregate_rigid_components

ROOT = Path(__file__).resolve().parents[2]
ROBOT = ROOT / 'robots/Goose_V0.1'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    relative_files = [
        'cad/exports/compute_module_mounts/manifest.json', 'cad/exports/compute_module_fixture/manifest.json',
        'cad/source/compute_module_fixture/assembly_scene.json', 'evidence/compute_module_mount_fit.json',
        'hardware/compute_module_mount_bom.json', 'evidence/compute_module_fixture_identity.json',
        'evidence/integrated_hardware_parameters.json', 'cad/source/integrated_hardware_candidate/scene.json']
    files = [ROBOT/p for p in relative_files]
    kit, fixture, fixture_scene, fit, bom, blend, ledger, scene = [json.loads(p.read_text()) for p in files]
    for record in [kit,fixture,fixture_scene,fit,bom,ledger,scene]:
        for relative,expected in record['source_hashes'].items():
            if sha(ROOT/relative)!=expected:
                raise ValueError(('Changed source chain',relative))
    for record in [kit,fixture]:
        for part in record['parts']:
            for output in part['files'].values():
                if sha(ROBOT/output['path'])!=output['sha256']:
                    raise ValueError(('Changed native output',part['name']))
    for relative,expected in blend['source_hashes'].items():
        if sha(ROBOT/relative)!=expected:
            raise ValueError(('Changed Blender source',relative))
    if blend['generator_sha256']!=sha(ROOT/'scripts/diagnostics/check_goose_mechanical_blend.py'):
        raise ValueError('Changed Blender identity validator')
    if (len(kit['parts']),len(fixture['parts']),len(fixture_scene['parts']))!=(42,2,50):
        raise ValueError('Declared scope incomplete')
    if not blend['all_source_identity_pass'] or blend['parts']!=50:
        raise ValueError('Fixture identity failed')
    if not fit['finite_screen_complete'] or not fit['valid_sources_sampled_fit_pass'] or fit['all_sources_sampled_fit_pass']:
        raise ValueError('Wrong finite fit status: invalid-source failure must remain explicit')
    invalid='radxa_zero_3w_invalid_vendor_solid_0'
    if fit['requested_poses']!=34 or fit['completed_poses']!=34 or fit['supplier_geometry_validity_pass']:
        raise ValueError('Incomplete screen or lost OEM source failure')
    for row in fit['cases']:
        if row['collisions'] or not row['complete'] or len(row['unresolved'])!=13 or any(
                invalid not in [pair['a'],pair['b']] for pair in row['unresolved']):
            raise ValueError(('Unexpected finite record',row['name']))
    if len(kit['nominal_thread_contacts'])!=12 or any(r['thread_engagement_mm']<2 for r in kit['nominal_thread_contacts']):
        raise ValueError('Thread stack incomplete')
    for baseline in fit['default_baseline']:
        if sha(ROOT/baseline['path'])!=baseline['sha256']:
            raise ValueError('Default changed')
    for supplier in fit['supplier_inputs']:
        if sha(ROOT/supplier['path'])!=supplier['sha256']:
            raise ValueError('Private source changed')
    for relative,expected in bom['outputs'].items():
        if sha(ROOT/relative)!=expected:
            raise ValueError('BOM output changed')
    with (ROBOT/'hardware/compute_module_mount_bom.csv').open(newline='') as file:
        rows=list(csv.DictReader(file))
    workbook=load_workbook(ROBOT/'hardware/compute_module_mount_bom.xlsx',read_only=True,data_only=True)
    sheet=workbook.active
    columns=[c.value for c in next(sheet.iter_rows())]
    excel_rows=[dict(zip(columns,row)) for row in sheet.iter_rows(min_row=2,values_only=True)]
    workbook.close()
    if len(rows)!=9 or len(excel_rows)!=9 or sum(int(r['quantity']) for r in rows)!=44:
        raise ValueError('BOM count mismatch')
    for actual,excel,expected in zip(rows,excel_rows,bom['rows']):
        for key,value in expected.items():
            cell=excel[key] if excel[key] is not None else ''
            matched=math.isclose(cell,value,rel_tol=1e-12,abs_tol=1e-14) if isinstance(value,float) and isinstance(cell,(float,int)) else cell==value
            if actual[key]!=str(value) or not matched:
                raise ValueError(('BOM discrepancy',expected['id'],key))
    if len(scene['parts'])!=419 or len({p['name'] for p in scene['parts']})!=419 or len(ledger['bodies'])!=19:
        raise ValueError('Integrated identity/count mismatch')
    if scene['nominal_conditional_mass_kg']!=ledger['nominal_conditional_mass_kg']:
        raise ValueError('Scene/ledger mass mismatch')
    native_geometry_refs=0
    for part in scene['parts']:
        if 'geometry_npz' in part:
            if sha(ROBOT/part['geometry_npz'])!=part['source_sha256']:
                raise ValueError(('Changed integrated geometry',part['name']))
            native_geometry_refs+=1
    if any(ledger[key] for key in ['training_release','hardware_freeze','full_assembly_pass','manufacturing_release','new_collision_model_generated']):
        raise ValueError('Premature integrated release')
    if not ledger['all_static_torques_pass'] or not ledger['middle_all_static_torques_pass'] or (
            ledger['static_contact_feasible'],ledger['middle_static_contact_feasible'])!=(63,63):
        raise ValueError('Declared clamp stations static screen failed')
    if len(ledger['front_50n_overload_failures'])!=63:
        raise ValueError('Front50N overload failure was lost')
    for body in ledger['bodies']:
        items=[i for i in ledger['items'] if i['body']==body['name']]
        expected=aggregate_rigid_components(items,ledger['pivots_world_at_zero_m'][body['name']])
        for key in ['mass_kg','com_local_m','inertia_at_com_body_kg_m2']:
            np.testing.assert_allclose(expected[key],body[key],atol=1e-12,rtol=1e-12)
    np.testing.assert_allclose(sum(b['mass_kg'] for b in ledger['bodies']),ledger['nominal_conditional_mass_kg'],atol=1e-12)
    with (ROBOT/'hardware/integrated_hardware_body_parameters.csv').open(newline='') as file:
        body_rows=list(csv.DictReader(file))
    if len(body_rows)!=19:
        raise ValueError('Integrated body table incomplete')
    for row,body in zip(body_rows,ledger['bodies']):
        if row['body']!=body['name']:
            raise ValueError('Body order mismatch')
        c,t=body['com_local_m'],body['inertia_at_com_body_kg_m2']
        expected=[body['mass_kg'],*c,t[0][0],t[1][1],t[2][2],t[0][1],t[0][2],t[1][2]]
        actual=[float(row[k]) for k in ['mass_kg','com_x_m','com_y_m','com_z_m','ixx','iyy','izz','ixy','ixz','iyz']]
        np.testing.assert_allclose(actual,expected,atol=1e-14,rtol=1e-12)
    docs=[ROBOT/'README.md',ROBOT/'design/compute_module_mount_checkpoint.md',ROBOT/'design/integrated_hardware_candidate.md']
    links=0
    for path in docs:
        for target in re.findall(r'\]\(([^)]+)\)',path.read_text()):
            if '://' in target or target.startswith('#'):
                continue
            links+=1
            if not (path.parent/target.split('#',1)[0]).exists():
                raise ValueError(('Missing local link',path,target))
    images=[]
    for view in ['three_quarter','rear','top']:
        path=ROBOT/'images/compute_module_mounts'/(view+'.png')
        with Image.open(path) as image:
            if image.size!=(1000,1000):
                raise ValueError('Unexpected render size')
            image.verify()
        images.append(dict(path=str(path.relative_to(ROOT)),sha256=sha(path),verified_png=True))
    helper_files=[ROOT/p for p in ['scripts/cad/build_goose_cad.py',
        'scripts/cad/build_goose_power_module_mounts.py','scripts/cad/goose_candidate_export.py',
        'scripts/cad/goose_nurbs_skin.py','scripts/cad/build_goose_actuator_interfaces.py',
        'scripts/models/build_goose_stage_two.py','scripts/models/build_goose_stage_one.py',
        'scripts/diagnostics/screen_goose_system_loads.py','src/sai_agent/goose/mass_properties.py']]
    inputs=files+docs+helper_files+[Path(__file__),ROOT/'scripts/models/render_goose_mechanical_preview.py',
                       ROBOT/'hardware/integrated_hardware_body_parameters.csv']
    result=dict(schema='goose_compute_mount_and_integrated_candidate_checkpoint_v1',integrity_pass=True,
        own_mount_parts=42, own_mount_quads=sum(r['quad_faces'] for r in fit['quad_sources']),
        fixture_parts=50, fixture_quads=blend['quad_faces'], fixture_source_identity_pass=True,
        finite_poses=34, valid_sources_sampled_fit_pass=True, all_sources_sampled_fit_pass=False,
        unresolved_source_pairs_per_pose=13, bom_rows=9,bom_pieces=44,csv_xlsx_match=True,
        compute_added_mass_kg=kit['conservative_added_mass_kg'],
        integrated_parts=419, integrated_hashed_npz_parts=native_geometry_refs,
        integrated_static_bodies=19, integrated_mass_kg=ledger['nominal_conditional_mass_kg'],
        front20n_static_cases_passed=63,middle50n_static_cases_passed=63,front50n_overload_failures_retained=63,
        closed_pose_mass_inertia_reaggregation_pass=True,integrated_body_csv_match=True,
        checked_local_links=links,missing_local_links=[],actual_fixture_renders=images,
        default_assembly_unchanged=True, manufacturing_release=False,full_assembly_pass=False,training_release=False,
        stage_three_four_complete=False,
        source_hashes={str(p.relative_to(ROOT)):sha(p) for p in inputs},
        next_main_scope='Interface/protection actual mounting and shell-to-frame support, mass/torque reserve, then fresh same-version dynamic model; no PPO expansion.')
    (ROBOT/'evidence/compute_module_mount_checkpoint.json').write_text(json.dumps(result,indent=2)+'\n')
    print('COMPUTE CHECKPOINT',42,'native;419integrated;',links,'links; fail states retained',flush=True)


if __name__=='__main__':
    main()
