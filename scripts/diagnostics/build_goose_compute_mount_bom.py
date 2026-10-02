"""Candidate compute-rack list with complete own-part coverage, no release."""
from pathlib import Path
import csv
import hashlib
import json

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
ROBOT = ROOT / 'robots/Goose_V0.1'


def main():
    cfg_path = ROBOT / 'configs/compute_module_mounts.json'
    kit_path = ROBOT / 'cad/exports/compute_module_mounts/manifest.json'
    fit_path = ROBOT / 'evidence/compute_module_mount_fit.json'
    cfg, kit, fit = [json.loads(p.read_text()) for p in [cfg_path, kit_path, fit_path]]
    parts = {p['name']: p for p in kit['parts']}
    rows, covered = [], set()
    for module in cfg['modules']:
        rows.append(dict(id=module['name'] + '_board', description=module['sku'], quantity=1,
            specification=module['vendor_revision'], mass_kg='',
            mass_basis='Bare individual masses not frozen; both remain inside existing54g combined reserve',
            step='', source=module['sources'][0], domestic_price_rmb='', purchase_release=False,
            installation=module['support_side']))
    for name, description in [('torso_compute_mount_chassis_plate', 'Modified3mm6061 chassis; replaces existing plate'),
                              ('compute_dual_shelf_frame_carrier', 'Original monolithic6061 compute rack; DFM pending')]:
        p = parts[name]
        covered.add(name)
        rows.append(dict(id=name, description=description, quantity=1,
            specification='Native STEP defines nominal geometry; tolerances/finish/local strength not released',
            mass_kg=p['mass_kg'], mass_basis='Native volume x density2700kg/m3', step=p['files']['step']['path'],
            source='configs/compute_module_mounts.json', domestic_price_rmb='', purchase_release=False,
            installation='Replacement chassis' if 'chassis' in name else 'Four dedicated M3x12 frame anchors'))
    groups = [
        ('pcb_nylon_washer', 'Nylon M2 washer', 'OD4/ID2.2/t0.5mm', lambda n: '_pcb_' in n and n.endswith('_washer')),
        ('pcb_m2x6', 'Steel M2x6 socket screw', 'HeadOD3.8/h2mm; M2 tapped supports', lambda n: '_pcb_' in n and n.endswith('_m2x6')),
        ('frame_m3_washer', 'Steel M3 washer', 'OD7/ID3.2/t0.5mm', lambda n: n.startswith('compute_frame_') and n.endswith('_washer')),
        ('frame_m3x12', 'Steel M3x12 socket screw', 'HeadOD5.5/h3mm; shaft12mm', lambda n: n.endswith('_m3x12')),
        ('frame_m3_nut', 'Steel M3 nut', 'AF5.5/h2.4mm', lambda n: n.endswith('_m3_nut'))]
    for identifier, description, specification, predicate in groups:
        chosen = [p for name, p in parts.items() if predicate(name)]
        if not chosen:
            raise ValueError('Empty hardware selection')
        covered.update(p['name'] for p in chosen)
        rows.append(dict(id=identifier, description=description, quantity=len(chosen), specification=specification,
            mass_kg=sum(p['mass_kg'] for p in chosen), mass_basis='Nominal native envelopes, not measured vendor hardware',
            step=chosen[0]['files']['step']['path'], source='configs/compute_module_mounts.json',
            domestic_price_rmb='', purchase_release=False, installation='Candidate nominal stack; locking/preload/tolerance pending'))
    if covered != set(parts):
        raise ValueError(('BOM coverage mismatch', set(parts)-covered))
    csv_path, xlsx_path = [ROBOT / 'hardware' / ('compute_module_mount_bom.' + x) for x in ['csv', 'xlsx']]
    columns = list(rows[0])
    with csv_path.open('w', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=columns, lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'compute_mount_candidate'
    sheet.append(columns)
    for row in rows:
        sheet.append([row[c] for c in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244B5A')
        sheet.column_dimensions[cell.column_letter].width = 12 if cell.value == 'quantity' else 44
    workbook.save(xlsx_path)
    inputs = [cfg_path, kit_path, fit_path, Path(__file__)]
    result = dict(schema='goose_compute_mount_incremental_bom_v1', rows=rows,
        native_part_coverage=len(covered), quantity_total=sum(r['quantity'] for r in rows),
        own_native_listed_mass_kg=kit['native_mass_kg'], conservative_robot_added_mass_kg=kit['conservative_added_mass_kg'],
        retained_compute_mass_reservation_kg=.054, includes_replacement_chassis=True,
        valid_sources_sampled_fit_pass=fit['valid_sources_sampled_fit_pass'],
        all_sources_sampled_fit_pass=fit['all_sources_sampled_fit_pass'],
        purchase_release=False, full_robot_bom=False,
        source_hashes={str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs},
        outputs={str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest() for p in [csv_path, xlsx_path]},
        limitations=['Radxa purchased revision/RAM/eMMC, individual module masses, domestic pricing, machining quote and exact fastener vendor unresolved.',
            'All source fit remains false because the original Radxa main solid is invalid. No electrical/thermal/wiring acceptance.',
            '42part listed mass includes a replacement chassis; only net29g is added, existing54g compute reserve retained.'])
    (ROBOT / 'hardware/compute_module_mount_bom.json').write_text(json.dumps(result, indent=2) + '\n')
    print('COMPUTE BOM', len(rows), 'rows', result['quantity_total'], 'pieces', flush=True)


if __name__ == '__main__':
    main()
