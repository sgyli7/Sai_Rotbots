"""Audit the462-piece shell-mount checkpoint without releasing the robot."""
from pathlib import Path
import csv
import hashlib
import json
import math
import re

import numpy as np
from openpyxl import load_workbook
from sai_agent.goose.mass_properties import aggregate_rigid_components

ROOT = Path(__file__).resolve().parents[2]
R = ROOT / 'robots/Goose_V0.1'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    paths = [R / p for p in ['cad/exports/torso_shell_mounts/manifest.json',
        'evidence/torso_shell_mount_parameters.json', 'cad/source/torso_shell_mount_fixture/assembly_scene.json',
        'evidence/torso_shell_mount_fit.json', 'hardware/torso_shell_mount_bom.json',
        'evidence/microduck_color_blocking_462_parts_render_manifest.json',
        'evidence/microduck_color_blocking_462_parts_pixel_validation.json',
        'cad/source/camera_closure_fixture/assembly_scene.json']]
    kit, params, scene, fit, bom, render, pixels, old = [json.loads(p.read_text()) for p in paths]
    for value in [kit, params, scene, fit, bom]:
        for relative, digest in value['source_hashes'].items():
            if sha(ROOT / relative) != digest:
                raise ValueError(('Changed source', relative))
    for part in kit['parts']:
        for file in part['files'].values():
            if sha(R / file['path']) != file['sha256']:
                raise ValueError(('Changed output', part['name']))
    before, after = [{p['name']: p for p in s['parts']} for s in [old, scene]]
    replaced = set(kit['replaces_existing_parts'])
    added = {p['name'] for p in kit['parts']}
    if len(after) != 462 or set(after) != set(before)-replaced | added or len(replaced) != 14 or len(added) != 30:
        raise ValueError('Incorrect replacement scope')
    for name in set(before)-replaced:
        if before[name] != after[name]:
            raise ValueError(('Unintended geometry change', name))
    for part in after.values():
        if part.get('geometry_npz') and sha(R / part['geometry_npz']) != part['source_sha256']:
            raise ValueError(('Changed scene geometry', part['name']))
    if not fit['incremental_geometry_pass'] or not fit['same_source_static_pass'] or fit['zero_only']:
        raise ValueError('Finite fit/static screen not passed')
    if fit['own_pair_failures'] or fit['own_pair_unresolved'] or len(fit['finite_cases']) != 18:
        raise ValueError('Unresolved shell-mount fit')
    if [len(fit[k]) for k in ['nominal_threads', 'retained_frame_nominal_threads', 'as_printed_heat_insert_interference']] != [6, 2, 4]:
        raise ValueError('Missing separately quantified installation contacts')
    for screen in fit['static_screens'].values():
        if screen['cases'] != 63 or screen['contact_feasible'] != 63 or not screen['all_static_torques_pass']:
            raise ValueError('Static case mismatch')
    for body in params['bodies']:
        actual = aggregate_rigid_components([p for p in params['items'] if p['body']==body['name']],
            params['pivots_world_at_zero_m'][body['name']])
        for key in ['mass_kg', 'com_local_m', 'inertia_at_com_body_kg_m2']:
            np.testing.assert_allclose(actual[key], body[key], rtol=1e-12, atol=1e-12)
        eigen = np.linalg.eigvalsh(np.array(body['inertia_at_com_body_kg_m2']))
        if eigen.min() <= 0 or eigen[-1] > eigen[:2].sum()+1e-12:
            raise ValueError(('Nonphysical inertia', body['name']))
    np.testing.assert_allclose(sum(b['mass_kg'] for b in params['bodies']), params['nominal_conditional_mass_kg'], atol=1e-12)
    old_params = json.loads((R/'evidence/camera_head_closure_parameters.json').read_text())
    for key in ['pivots_world_at_zero_m', 'neutral_joint_order', 'rigid_coordinate_lift_m']:
        if params[key] != old_params[key]:
            raise ValueError('Changed mechanical coordinate contract')
    if len(params['bodies']) != 19 or params['active_axes'] != 18 or not bom['original30g_unclosed_shell_hardware_reserve_retained']:
        raise ValueError('Missing parameter scope or mass reserve')
    for relative, digest in bom['outputs'].items():
        if sha(R/relative) != digest:
            raise ValueError('Changed BOM output')
    with (R/'hardware/torso_shell_mount_bom.csv').open(newline='') as file:
        rows = list(csv.DictReader(file))
    book = load_workbook(R/'hardware/torso_shell_mount_bom.xlsx',read_only=True,data_only=True)
    sheet = book.active
    fields = [c.value for c in next(sheet.iter_rows())]
    excel = [dict(zip(fields,row)) for row in sheet.iter_rows(min_row=2,values_only=True)]
    book.close()
    if len(rows) != 30 or len(excel) != 30 or sum(p['quantity'] for p in bom['rows']) != 30:
        raise ValueError('Wrong BOM piece count')
    for a,b,expected in zip(rows,excel,bom['rows']):
        for key,value in expected.items():
            cell = b[key] if b[key] is not None else ''
            matches = math.isclose(cell,value,rel_tol=1e-12,abs_tol=1e-14) if isinstance(value,float) else cell==value
            if a[key] != str(value) or not matches:
                raise ValueError(('BOM CSV/XLSX mismatch',key))
    if render['source_scene_sha256'] != sha(paths[2]) or pixels['render_manifest_sha256'] != sha(paths[5]):
        raise ValueError('Mixed image/model source')
    if not render['identical_geometry_verified'] or render['geometry_modifiers_created'] or len(pixels['images']) != 5:
        raise ValueError('Unverified render geometry')
    for row in pixels['images']:
        if not row['actual_pixels_inspected'] or sha(R/'images/microduck_color_blocking_462_parts'/row['file']) != row['sha256']:
            raise ValueError('Unverified actual image')
    documents = [R/'README.md',R/'design/torso_shell_mount_checkpoint.md',
        R/'design/microduck_color_blocking_462_parts.md',R/'design/hardware_milestones.md',R/'hardware/power_release_checkpoint.md']
    links = 0
    report_path = R/'evidence/torso_shell_mount_checkpoint.json'
    for path in documents:
        for target in re.findall(r'\]\(([^)]+)\)',path.read_text()):
            if '://' in target or target.startswith('#'):
                continue
            resolved = (path.parent/target.split('#',1)[0]).resolve()
            if not resolved.exists() and resolved != report_path:
                raise ValueError(('Broken relative link',path,target))
            links += 1
    for value in [params,fit]:
        if any(value[k] for k in ['manufacturing_release','full_assembly_pass','final_appearance_pass','training_release']):
            raise ValueError('Premature robot release')
    result = dict(schema='goose_torso_shell_mount_checkpoint_v1',integrity_pass=True,
        source_objects=462,native_increment_parts=30,unchanged_predecessor_objects=432,
        active_axes=18,closed_pose_bodies=19,mass_kg=params['nominal_conditional_mass_kg'],
        finite_poses_passed=18,front20n_static_cases_passed=63,middle50n_static_cases_passed=63,
        actual_same_source_images=5,bom_csv_xlsx_json_match=True,local_links_checked=links,
        stage_three_four_complete=False,manufacturing_release=False,full_assembly_pass=False,
        appearance_accepted_by_user=False,training_release=False,
        source_hashes={str(p.relative_to(ROOT)):sha(p) for p in paths+documents+[Path(__file__)]})
    report_path.write_text(json.dumps(result,indent=2)+'\n')
    print('SHELL CHECKPOINT',result['mass_kg'],'kg,462objects,18poses,5images,',links,'links')


if __name__ == '__main__':
    main()
