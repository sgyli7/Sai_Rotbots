"""Audit source, output and document identity of this finite mounting delivery."""
from pathlib import Path
import csv
import hashlib
import json
import math
import re

from PIL import Image
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
ROBOT = ROOT / 'robots/Goose_V0.1'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    files = [ROBOT / relative for relative in [
        'cad/exports/power_module_mounts/manifest.json',
        'cad/exports/power_module_fixture/manifest.json',
        'cad/source/power_module_fixture/assembly_scene.json',
        'evidence/power_module_mount_fit.json',
        'hardware/power_module_mount_bom.json',
        'evidence/power_module_fixture_identity.json']]
    kit, reference, scene, fit, bom, blend = [json.loads(path.read_text()) for path in files]
    for record in [kit, reference, scene, fit, bom]:
        for relative, expected in record['source_hashes'].items():
            if sha(ROOT / relative) != expected:
                raise ValueError(('Stale source chain', relative))
    for record in [kit, reference]:
        for part in record['parts']:
            for output in part['files'].values():
                if sha(ROBOT / output['path']) != output['sha256']:
                    raise ValueError(('Changed native output', part['name']))
    for relative, expected in blend['source_hashes'].items():
        if sha(ROBOT / relative) != expected:
            raise ValueError(('Changed rendered source', relative))
    validator = ROOT / 'scripts/diagnostics/check_goose_mechanical_blend.py'
    if sha(validator) != blend['generator_sha256']:
        raise ValueError('Changed Blender identity checker')
    if (len(kit['parts']), len(reference['parts']), len(scene['parts'])) != (36, 2, 43):
        raise ValueError('Incomplete declared part scope')
    if not blend['all_source_identity_pass'] or blend['parts'] != 43:
        raise ValueError('Saved Blender source identity failed')
    if not fit['named_sampled_fit_pass'] or not fit['bare_modules_in_original_boxes']:
        raise ValueError('Finite installation screen failed')
    if fit['requested_poses'] != 34 or fit['completed_poses'] != 34 or any(
            row['collisions'] or row['unresolved'] or not row['complete'] for row in fit['cases']):
        raise ValueError('Incomplete finite native fit')
    if len(kit['nominal_thread_contacts']) != 10 or any(
            row['thread_engagement_mm'] < 2 for row in kit['nominal_thread_contacts']):
        raise ValueError('Unexpected thread scope or insufficient nominal engagement')
    for baseline in fit['default_baseline']:
        if sha(ROOT / baseline['path']) != baseline['sha256']:
            raise ValueError(('Default assembly changed', baseline['path']))
    for supplier in fit['supplier_inputs']:
        if sha(ROOT / supplier['path']) != supplier['sha256']:
            raise ValueError(('Changed private OEM source', supplier['name']))
    for relative, expected in bom['outputs'].items():
        if sha(ROOT / relative) != expected:
            raise ValueError(('Changed BOM output', relative))
    csv_path = ROBOT / 'hardware/power_module_mount_bom.csv'
    with csv_path.open(newline='') as file:
        rows = list(csv.DictReader(file))
    workbook = load_workbook(ROBOT / 'hardware/power_module_mount_bom.xlsx', read_only=True, data_only=True)
    sheet = workbook.active
    columns = [cell.value for cell in next(sheet.iter_rows())]
    excel_rows = [dict(zip(columns, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    workbook.close()
    expected = bom['rows']
    if len(rows) != len(excel_rows) or len(rows) != len(expected) or sum(int(r['quantity']) for r in rows) != 38:
        raise ValueError('BOM row or quantity mismatch')
    for a, b, e in zip(rows, excel_rows, expected):
        for key, value in e.items():
            excel_value = b[key] if b[key] is not None else ''
            excel_matches = (math.isclose(excel_value, value, rel_tol=1e-12, abs_tol=1e-14)
                             if isinstance(value, float) and isinstance(excel_value, (float, int))
                             else excel_value == value)
            if a[key] != str(value) or not excel_matches:
                raise ValueError(('CSV/XLSX/manifest discrepancy', e['id'], key))

    docs = [ROBOT / 'README.md', ROBOT / 'design/power_module_mount_checkpoint.md']
    links = 0
    for path in docs:
        for target in re.findall(r'\]\(([^)]+)\)', path.read_text()):
            if '://' in target or target.startswith('#'):
                continue
            links += 1
            if not (path.parent / target.split('#', 1)[0]).exists():
                raise ValueError(('Missing local link', path, target))
    images = []
    for view in ['three_quarter', 'top']:
        path = ROBOT / 'images/power_module_mounts' / (view + '.png')
        with Image.open(path) as image:
            if image.size != (1000, 1000):
                raise ValueError('Unexpected render dimensions')
            image.verify()
        images.append(dict(path=str(path.relative_to(ROOT)), sha256=sha(path), verified_png=True))
    inputs = files + docs + [validator, Path(__file__), ROOT / 'scripts/models/render_goose_mechanical_preview.py']
    result = dict(schema='goose_power_mount_checkpoint_v1', integrity_pass=True,
        native_parts=36, native_quad_faces=sum(row['faces'] for row in fit['quad_sources']),
        fixture_parts=43, fixture_quads=blend['quad_faces'], saved_blend_identity_pass=True,
        finite_poses=34, finite_named_fit_pass=True, bare_modules_in_original_boxes=True,
        incremental_bom_rows=len(rows), incremental_bom_items=38, csv_xlsx_match=True,
        conservative_added_mass_kg=fit['conservative_candidate_added_mass_kg'],
        checked_local_links=links, missing_local_links=[], actual_fixture_renders=images,
        appearance_approval=False,
        default_assembly_unchanged=True, installed=False, manufacturing_release=False,
        electrical_release=False, full_assembly_pass=False, stage_three_four_complete=False,
        source_hashes={str(path.relative_to(ROOT)): sha(path) for path in inputs},
        next_main_scope='Compute/interface module mounting and shell-to-frame support, then one combined assembly revision; no PPO extension.')
    output = ROBOT / 'evidence/power_module_mount_checkpoint.json'
    output.write_text(json.dumps(result, indent=2) + '\n')
    print('POWER CHECKPOINT', result['native_parts'], 'parts', result['native_quad_faces'], 'quads',
          result['finite_poses'], 'poses', links, 'local links', flush=True)


if __name__ == '__main__':
    main()
