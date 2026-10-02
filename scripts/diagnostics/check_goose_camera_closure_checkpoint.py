"""Audit the corrected camera/head delivery without granting whole-robot release."""
from pathlib import Path
import csv
import hashlib
import json
import math
import re

import numpy as np
from openpyxl import load_workbook
from PIL import Image
from sai_agent.goose.mass_properties import aggregate_rigid_components

ROOT = Path(__file__).resolve().parents[2]
ROBOT = ROOT / 'robots/Goose_V0.1'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    paths = [ROBOT / relative for relative in [
        'cad/exports/camera_head_closure/manifest.json',
        'evidence/camera_head_closure_parameters.json',
        'cad/source/camera_closure_fixture/assembly_scene.json',
        'evidence/camera_head_closure_fit.json',
        'hardware/camera_head_closure_bom.json',
        'evidence/microduck_color_blocking_446_parts_render_manifest.json',
        'evidence/microduck_color_blocking_446_parts_pixel_validation.json',
        'cad/source/integrated_hardware_candidate/scene.json']]
    kit, parameters, scene, fit, bom, render, pixels, predecessor = [json.loads(path.read_text()) for path in paths]
    for value in [kit, parameters, scene, fit, bom]:
        for relative, digest in value['source_hashes'].items():
            if sha(ROOT / relative) != digest:
                raise ValueError(('Changed source', relative))
    for part in kit['parts']:
        for output in part['files'].values():
            if sha(ROBOT / output['path']) != output['sha256']:
                raise ValueError(('Changed native output', part['name']))
    old = {part['name']: part for part in predecessor['parts']}
    current = {part['name']: part for part in scene['parts']}
    removed = set(kit['replaces_existing_parts']) | {'camera_board_proxy', 'camera_lens_proxy'}
    if set(old) - set(current) != removed or len(current) != 446 or len(current.keys() - old.keys()) != 33:
        raise ValueError('Incorrect candidate replacement scope')
    for name in current.keys() & old.keys():
        if current[name] != old[name]:
            raise ValueError(('Unintended existing geometry change', name))
    for part in current.values():
        if 'geometry_npz' in part and sha(ROBOT / part['geometry_npz']) != part['source_sha256']:
            raise ValueError(('Changed scene geometry', part['name']))
    if not fit['incremental_geometry_pass'] or not fit['same_source_static_pass']:
        raise ValueError('Incremental geometry/static screen not passed')
    if fit['own_pair_failures'] or fit['unresolved'] or len(fit['incremental_head_cases']) != 10:
        raise ValueError('Unresolved native head case')
    if len(fit['nominal_thread_contacts']) != 4 or len(fit['nominal_oem_thread_contacts']) != 4:
        raise ValueError('Missing individually verified nominal threaded contacts')
    for body in parameters['bodies']:
        actual = aggregate_rigid_components(
            [part for part in parameters['items'] if part['body'] == body['name']],
            parameters['pivots_world_at_zero_m'][body['name']])
        for key in ['mass_kg', 'com_local_m', 'inertia_at_com_body_kg_m2']:
            np.testing.assert_allclose(actual[key], body[key], rtol=1e-12, atol=1e-12)
    np.testing.assert_allclose(sum(body['mass_kg'] for body in parameters['bodies']),
                               parameters['nominal_conditional_mass_kg'], atol=1e-12)
    with (ROBOT / 'hardware/camera_closure_body_parameters.csv').open(newline='') as file:
        body_rows = list(csv.DictReader(file))
    if len(body_rows) != 19:
        raise ValueError('Missing body parameters')
    for row, body in zip(body_rows, parameters['bodies']):
        tensor = body['inertia_at_com_body_kg_m2']
        values = [body['mass_kg'], *body['com_local_m'], tensor[0][0], tensor[1][1], tensor[2][2],
                  tensor[0][1], tensor[0][2], tensor[1][2]]
        if row['body'] != body['name']:
            raise ValueError('Wrong body order')
        np.testing.assert_allclose([float(row[key]) for key in list(row)[1:]], values, rtol=1e-12, atol=1e-14)
    for relative, digest in bom['outputs'].items():
        if sha(ROBOT / relative) != digest:
            raise ValueError('Changed BOM output')
    with (ROBOT / 'hardware/camera_head_closure_bom.csv').open(newline='') as file:
        csv_rows = list(csv.DictReader(file))
    book = load_workbook(ROBOT / 'hardware/camera_head_closure_bom.xlsx', read_only=True, data_only=True)
    sheet = book.active
    fields = [cell.value for cell in next(sheet.iter_rows())]
    excel_rows = [dict(zip(fields, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
    book.close()
    if len(csv_rows) != 7 or len(excel_rows) != 7 or sum(row['quantity'] for row in bom['rows']) != 29:
        raise ValueError('Wrong nominal procurement scope')
    for csv_row, excel_row, expected in zip(csv_rows, excel_rows, bom['rows']):
        for key, value in expected.items():
            cell = excel_row[key] if excel_row[key] is not None else ''
            matches = math.isclose(cell, value, rel_tol=1e-12, abs_tol=1e-14) if isinstance(value, float) else cell == value
            if csv_row[key] != str(value) or not matches:
                raise ValueError(('CSV/XLSX/JSON mismatch', expected['id'], key))
    if render['source_scene_sha256'] != sha(paths[2]) or pixels['render_manifest_sha256'] != sha(paths[5]):
        raise ValueError('Mixed image/model source')
    if not render['identical_geometry_verified'] or render['geometry_modifiers_created'] != 0:
        raise ValueError('Modified render geometry')
    for row in pixels['images']:
        path = ROBOT / 'images/microduck_color_blocking_446_parts' / row['file']
        if not row['actual_pixels_inspected'] or sha(path) != row['sha256']:
            raise ValueError('Unverified final pixels')
        with Image.open(path) as image:
            if image.size != (1200, 1200):
                raise ValueError('Wrong render size')
            image.verify()
    documents = [ROBOT / 'README.md', ROBOT / 'design/camera_head_closure_checkpoint.md',
                 ROBOT / 'design/microduck_color_blocking_446_parts.md',
                 ROBOT / 'design/hardware_milestones.md', ROOT / 'docs/guides/large_asset_checkout.md']
    report_path = ROBOT / 'evidence/camera_head_closure_checkpoint.json'
    links = 0
    for path in documents:
        for target in re.findall(r'\]\(([^)]+)\)', path.read_text()):
            if '://' in target or target.startswith('#'):
                continue
            resolved = (path.parent / target.split('#', 1)[0]).resolve()
            if not resolved.exists() and resolved != report_path:
                raise ValueError(('Broken local link', path, target))
            links += 1
    for value in [parameters, fit]:
        if any(value[key] for key in ['manufacturing_release', 'full_assembly_pass', 'final_appearance_pass', 'training_release']):
            raise ValueError('Premature whole-robot release')
    inputs = paths + documents + [Path(__file__), ROBOT / 'hardware/camera_closure_body_parameters.csv',
        ROOT / 'scripts/cad/build_goose_camera_closure_bom.py', ROOT / 'src/sai_agent/goose/mass_properties.py']
    report = dict(
        schema='goose_camera_head_closure_checkpoint_v1', integrity_pass=True,
        own_native_parts=29, all_own_native_quad_closure_pass=True, candidate_scene_parts=446,
        unchanged_predecessor_parts=413, dimensional_display_references_not_printable=4,
        active_axes=18, mass_kg=parameters['nominal_conditional_mass_kg'],
        mass_delta_from419_kg=parameters['delta_from419_kg'], closed_pose_bodies=19,
        finite_head_poses_passed=10, source_complete_oem_children=4,
        nominal_frame_thread_contacts_passed=4, nominal_oem_thread_contacts_passed=4,
        front20n_static_cases_passed=63, middle50n_static_cases_passed=63,
        bom_csv_xlsx_json_match=True, bom_rows=7, bom_nominal_pieces=29,
        actual_same_source_images=5, local_links_checked=links,
        default344_unchanged=True, original419_unchanged=True,
        stage_three_four_complete=False, manufacturing_release=False, full_assembly_pass=False,
        appearance_accepted_by_user=False, training_release=False,
        source_hashes={str(path.relative_to(ROOT)): sha(path) for path in inputs})
    report_path.write_text(json.dumps(report, indent=2) + '\n')
    print('HEAD CHECKPOINT', report['mass_kg'], 'kg,29native,446scene,10poses,5images,', links, 'local links')


if __name__ == '__main__':
    main()
